"""Exercise cached Excel formula values without application startup services."""

import ast
import io
from pathlib import Path
import re
from typing import Any, Dict, List, Optional
import unicodedata
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook


def _attendance_parser():
    """Load the actual pure parser functions, avoiding Firebase initialization."""
    source = Path(__file__).resolve().parents[1] / "src" / "app.py"
    names = {
        "clean_feedback_text", "hours_to_minutes", "truthy_attendance_flag",
        "normalize_attendance_key", "is_attendance_header_row",
        "attendance_rows_from_matrix", "row_value", "parse_attendance_xlsx_bytes",
        "aggregate_attendance_rows", "ATTENDANCE_HOURS_HEADER_PREFIXES",
        "ATTENDANCE_HOURS_HEADER_KEYS", "ATTENDANCE_SUMMARY_ROW_KEYS",
    }
    nodes = [
        node for node in ast.parse(source.read_text(encoding="utf-8")).body
        if (isinstance(node, ast.FunctionDef) and node.name in names)
        or (isinstance(node, ast.Assign) and any(
            isinstance(target, ast.Name) and target.id in names for target in node.targets
        ))
    ]
    namespace = {
        "io": io, "re": re, "unicodedata": unicodedata,
        "Any": Any, "Dict": Dict, "List": List, "Optional": Optional,
    }
    exec(compile(ast.Module(body=nodes, type_ignores=[]), str(source), "exec"), namespace)
    assert names <= namespace.keys(), f"Missing attendance definitions: {sorted(names - namespace.keys())}"
    return namespace["parse_attendance_xlsx_bytes"]


def test_attendance_xlsx_uses_cached_formula_hours():
    workbook = Workbook()
    workbook.active.append(["date", "work_hours"])
    workbook.active.append(["2026-09-12", "=8+0.5"])
    original = io.BytesIO()
    workbook.save(original)
    workbook.close()

    # openpyxl writes formulas without calculating them. Supply the cached value
    # that Excel saves, keeping the formula and all workbook data in memory.
    cached = io.BytesIO()
    with ZipFile(io.BytesIO(original.getvalue())) as source, ZipFile(cached, "w") as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                formula = b"<f>8+0.5</f><v></v>"
                assert content.count(formula) == 1
                content = content.replace(formula, b"<f>8+0.5</f><v>8.5</v>", 1)
            target.writestr(entry, content)

    payload = cached.getvalue()
    formula_workbook = load_workbook(io.BytesIO(payload), read_only=True)
    try:
        assert formula_workbook.active["B2"].value == "=8+0.5"
    finally:
        formula_workbook.close()

    result = _attendance_parser()(payload)
    assert result["parsed_rows"] == 1
    assert result["work_minutes"] == 510  # The cached 8.5 hours, not the formula text.
