"""Exercise the Pillow APIs used by the demo and editable PPTX pipeline."""

from io import BytesIO

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from scripts import generate_seedance_demo_video as demo


def test_procedural_demo_frame_survives_png_roundtrip():
    # One frame only: no video encoder, filesystem output, or external service.
    frame = demo.draw_frame(demo.FRAME_COUNT // 2)
    assert frame.mode == "RGB"
    assert frame.size == (1280, 720)
    assert any(low != high for low, high in frame.getextrema())
    image_bytes = BytesIO()
    frame.save(image_bytes, format="PNG")
    image_bytes.seek(0)
    with Image.open(image_bytes) as restored:
        restored.load()
        assert restored.format == "PNG"
        assert restored.size == frame.size
        assert restored.getpixel((320, 240)) == frame.getpixel((320, 240))


def test_pillow_image_and_editable_text_survive_pptx_roundtrip():
    image_bytes = BytesIO()
    Image.new("RGB", (64, 36), "#165dff").save(image_bytes, format="PNG")
    image_bytes.seek(0)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(image_bytes, Inches(1), Inches(1), width=Inches(2))
    slide.shapes.add_textbox(Inches(1), Inches(3), Inches(5), Inches(1)).text = (
        "Mighty Skill-Bridge — 編集可能テキスト"
    )
    deck_bytes = BytesIO()
    deck.save(deck_bytes)
    deck_bytes.seek(0)
    restored = Presentation(deck_bytes)
    assert len(restored.slides) == 1
    shapes = restored.slides[0].shapes
    pictures = [shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1
    assert pictures[0].image.size == (64, 36)
    assert pictures[0].image.content_type == "image/png"
    assert [shape.text for shape in shapes if shape.has_text_frame] == [
        "Mighty Skill-Bridge — 編集可能テキスト"
    ]


def test_demo_font_fallback_can_measure_text(monkeypatch):
    monkeypatch.setattr(demo.Path, "exists", lambda _: False)
    font = demo.load_font(18)
    bounds = font.getbbox("Mighty Skill-Bridge")
    assert bounds[2] > bounds[0]
    assert bounds[3] > bounds[1]
